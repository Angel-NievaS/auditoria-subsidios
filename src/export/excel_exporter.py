"""Genera SUBSIDIOS_[CARPETA]_[DDMMYYYY].xlsx.
Estructura: hoja Resumen (primera) + una hoja por alumno.
"""
import os
import re
from itertools import groupby as _gb

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.domain.models import Comprobante, ResultadoAlumno, ResultadoAuditoria

# ── Columnas de la hoja por alumno ───────────────────────────────────────────
_HDRS_ALUMNO = [
    "Semana", "Pago", "Tipo doc",
    "Monto", "RUT", "Tipo cuenta", "N° cuenta", "Banco",
    "Fecha pago", "Estado", "Alertas",
]
_WIDTHS_ALUMNO = [8, 8, 18, 14, 16, 16, 14, 16, 12, 12, 60]

# ── Columnas del Resumen ──────────────────────────────────────────────────────
_HDRS_RESUMEN = [
    "Curso", "Alumno",
    "Sem.\nSubsidio", "Sem.\nCuidados",
    "Total Subsidio", "Total Cuidados", "Total General",
    "Alertas",
]
_WIDTHS_RESUMEN = [25, 42, 12, 12, 16, 16, 16, 10]

# ── Estilos ───────────────────────────────────────────────────────────────────
_FILL = {
    "header":    PatternFill("solid", fgColor="1F4E79"),
    "curso":     PatternFill("solid", fgColor="2C5F8A"),
    "cuidados":  PatternFill("solid", fgColor="2E7D32"),
    "subtotal":  PatternFill("solid", fgColor="DDEEFF"),
    "total_sub": PatternFill("solid", fgColor="1F4E79"),
    "total_cui": PatternFill("solid", fgColor="2E7D32"),
    "error":     PatternFill("solid", fgColor="FFCCCC"),
    "warn":      PatternFill("solid", fgColor="FFFFCC"),
    "back":      PatternFill("solid", fgColor="E8EDF5"),
    "alumno_bg": PatternFill("solid", fgColor="EEF3FB"),
    "alerta_r":  PatternFill("solid", fgColor="FFCCCC"),
}
_HDR_FONT  = Font(bold=True, color="FFFFFF")
_SUB_FONT  = Font(bold=True, italic=True)
_TOT_FONT  = Font(bold=True, color="FFFFFF")
_LINK_FONT = Font(color="0563C1", underline="single", bold=True)
_ALIGN_C   = Alignment(horizontal="center", vertical="center", wrap_text=True)
_ALIGN_L   = Alignment(horizontal="left",   vertical="center", indent=1)
_ALIGN_R   = Alignment(horizontal="right",  vertical="center")


def export_to_excel(resultado: ResultadoAuditoria, output_dir: str) -> str:
    nombre = _nombre_archivo(resultado)
    ruta = os.path.join(output_dir, nombre)

    wb = Workbook()

    # Mapear alumno.nombre → nombre de hoja único (≤31 chars)
    usados: set[str] = {"Resumen"}
    nombres_hojas: dict[str, str] = {}
    for alumno in resultado.alumnos:
        nombres_hojas[alumno.nombre] = _nombre_hoja(alumno.nombre, usados)

    _build_resumen(wb, resultado, nombres_hojas)
    for alumno in resultado.alumnos:
        _build_alumno_sheet(wb, alumno, nombres_hojas[alumno.nombre])

    wb.save(ruta)
    return ruta


# ── Hoja Resumen ──────────────────────────────────────────────────────────────

def _build_resumen(
    wb: Workbook,
    resultado: ResultadoAuditoria,
    nombres_hojas: dict[str, str],
) -> None:
    ws = wb.active
    ws.title = "Resumen"

    # Encabezado
    ws.append(_HDRS_RESUMEN)
    for col in range(1, len(_HDRS_RESUMEN) + 1):
        c = ws.cell(1, col)
        c.font = _HDR_FONT
        c.fill = _FILL["header"]
        c.alignment = _ALIGN_C
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"

    alumnos_sorted = sorted(resultado.alumnos, key=lambda a: a.curso)

    for curso, grupo_iter in _gb(alumnos_sorted, key=lambda a: a.curso):
        grupo = list(grupo_iter)
        label_curso = curso if curso else "Sin código de curso"

        # Encabezado de curso
        ws.append([label_curso] + [""] * (len(_HDRS_RESUMEN) - 1))
        fila_cur = ws.max_row
        for col in range(1, len(_HDRS_RESUMEN) + 1):
            ws.cell(fila_cur, col).fill = _FILL["curso"]
            ws.cell(fila_cur, col).font = Font(bold=True, color="FFFFFF", size=11)
        ws.cell(fila_cur, 1).alignment = _ALIGN_L
        ws.row_dimensions[fila_cur].height = 18

        # Una fila por alumno
        for alumno in grupo:
            sem_sub = len({c.semana for c in alumno.comprobantes
                           if c.categoria == "subsidio" and c.semana is not None})
            sem_cui = len({c.semana for c in alumno.comprobantes
                           if c.categoria == "cuidados" and c.semana is not None})
            tot_sub = sum(c.monto for c in alumno.comprobantes
                          if c.categoria == "subsidio" and c.monto)
            tot_cui = sum(c.monto for c in alumno.comprobantes
                          if c.categoria == "cuidados" and c.monto)
            errores, avisos = _contar_alertas(alumno)
            alertas_txt = _formato_alertas(errores, avisos)

            ws.append([
                alumno.curso,
                alumno.nombre,
                sem_sub or "",
                sem_cui or "",
                tot_sub,
                tot_cui,
                tot_sub + tot_cui,
                alertas_txt,
            ])
            fila = ws.max_row

            # Nombre → hyperlink a la hoja del alumno
            hoja = nombres_hojas[alumno.nombre]
            cell_n = ws.cell(fila, 2)
            cell_n.hyperlink = f"#{_quote_sheet(hoja)}!A1"
            cell_n.font = _LINK_FONT
            cell_n.alignment = _ALIGN_L

            for col in (5, 6, 7):
                ws.cell(fila, col).number_format = "#,##0"
                ws.cell(fila, col).alignment = _ALIGN_R
            for col in (3, 4):
                ws.cell(fila, col).alignment = _ALIGN_C

            if errores or avisos:
                cell_a = ws.cell(fila, 8)
                cell_a.fill = _FILL["alerta_r"] if errores else _FILL["warn"]
                cell_a.font = Font(bold=True, color="8B0000" if errores else "7B6000")
                cell_a.alignment = _ALIGN_C

        # Subtotal del curso
        tot_sub_c = sum(c.monto for a in grupo for c in a.comprobantes
                        if c.categoria == "subsidio" and c.monto)
        tot_cui_c = sum(c.monto for a in grupo for c in a.comprobantes
                        if c.categoria == "cuidados" and c.monto)
        ws.append(["", f"Subtotal — {len(grupo)} alumno{'s' if len(grupo) != 1 else ''}",
                   "", "", tot_sub_c, tot_cui_c, tot_sub_c + tot_cui_c, ""])
        fila_sub = ws.max_row
        for col in range(1, len(_HDRS_RESUMEN) + 1):
            ws.cell(fila_sub, col).fill = _FILL["subtotal"]
            ws.cell(fila_sub, col).font = _SUB_FONT
        for col in (5, 6, 7):
            ws.cell(fila_sub, col).number_format = "#,##0"
            ws.cell(fila_sub, col).alignment = _ALIGN_R
        ws.row_dimensions[fila_sub].height = 15

    # Fila TOTAL GENERAL
    tot_sub_all = sum(c.monto for a in resultado.alumnos for c in a.comprobantes
                      if c.categoria == "subsidio" and c.monto)
    tot_cui_all = sum(c.monto for a in resultado.alumnos for c in a.comprobantes
                      if c.categoria == "cuidados" and c.monto)
    err_all  = sum(e for e, _ in (_contar_alertas(a) for a in resultado.alumnos))
    avis_all = sum(v for _, v in (_contar_alertas(a) for a in resultado.alumnos))
    ws.append([
        "TOTAL GENERAL",
        f"{len(resultado.alumnos)} alumnos",
        "", "",
        tot_sub_all, tot_cui_all, tot_sub_all + tot_cui_all,
        _formato_alertas(err_all, avis_all),
    ])
    fila_tot = ws.max_row
    for col in range(1, len(_HDRS_RESUMEN) + 1):
        ws.cell(fila_tot, col).fill = _FILL["total_sub"]
        ws.cell(fila_tot, col).font = _TOT_FONT
    for col in (5, 6, 7):
        ws.cell(fila_tot, col).number_format = "#,##0"
        ws.cell(fila_tot, col).alignment = _ALIGN_R
    if err_all or avis_all:
        ws.cell(fila_tot, 8).fill = _FILL["alerta_r"] if err_all else _FILL["warn"]
    ws.cell(fila_tot, 8).alignment = _ALIGN_C
    ws.row_dimensions[fila_tot].height = 20

    for i, w in enumerate(_WIDTHS_RESUMEN, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ── Hoja por alumno ───────────────────────────────────────────────────────────

def _build_alumno_sheet(
    wb: Workbook,
    alumno: ResultadoAlumno,
    nombre_hoja: str,
) -> None:
    ws = wb.create_sheet(nombre_hoja)
    ncols = len(_HDRS_ALUMNO)
    last_col = get_column_letter(ncols)

    # Fila 1: Botón "← Volver al Resumen"
    ws.append(["← Volver al Resumen"])
    fila = ws.max_row
    ws.merge_cells(f"A{fila}:{last_col}{fila}")
    cell = ws.cell(fila, 1)
    cell.hyperlink = "#Resumen!A1"
    cell.font = _LINK_FONT
    cell.fill = _FILL["back"]
    cell.alignment = _ALIGN_L
    ws.row_dimensions[fila].height = 18

    # Fila 2: Nombre y curso
    titulo = f"{alumno.nombre}  —  {alumno.curso}" if alumno.curso else alumno.nombre
    ws.append([titulo])
    fila = ws.max_row
    ws.merge_cells(f"A{fila}:{last_col}{fila}")
    cell = ws.cell(fila, 1)
    cell.font = Font(bold=True, size=12, color="1A3A5C")
    cell.fill = _FILL["alumno_bg"]
    cell.alignment = _ALIGN_L
    ws.row_dimensions[fila].height = 24

    # Alertas del alumno (solo error/warn)
    for alerta in alumno.alertas:
        if alerta.severidad == "info":
            continue
        ws.append([f"{alerta.icono} {alerta.mensaje}"])
        fila = ws.max_row
        ws.merge_cells(f"A{fila}:{last_col}{fila}")
        fgColor = "FFCCCC" if alerta.severidad == "error" else "FFFFCC"
        ws.cell(fila, 1).fill = PatternFill("solid", fgColor=fgColor)
        ws.cell(fila, 1).font = Font(size=10)
        ws.cell(fila, 1).alignment = _ALIGN_L

    # Una tabla por categoría
    for cat, cat_label, fill_key in (
        ("subsidio", "SUBSIDIO", "header"),
        ("cuidados", "CUIDADOS", "cuidados"),
    ):
        comps_cat = [c for c in alumno.comprobantes if c.categoria == cat]
        if not comps_cat:
            continue

        # Fila vacía + título categoría
        ws.append([])
        ws.append([f"  {cat_label}"])
        fila = ws.max_row
        ws.merge_cells(f"A{fila}:{last_col}{fila}")
        ws.cell(fila, 1).fill = _FILL[fill_key]
        ws.cell(fila, 1).font = Font(bold=True, color="FFFFFF", size=11)
        ws.cell(fila, 1).alignment = _ALIGN_L
        ws.row_dimensions[fila].height = 18

        # Encabezado de tabla
        ws.append(_HDRS_ALUMNO)
        fila = ws.max_row
        for col in range(1, ncols + 1):
            c = ws.cell(fila, col)
            c.font = _HDR_FONT
            c.fill = _FILL["header"]
            c.alignment = _ALIGN_C
        ws.row_dimensions[fila].height = 20

        # Datos agrupados por semana
        por_semana: dict[int | None, list[Comprobante]] = {}
        for c in comps_cat:
            por_semana.setdefault(c.semana, []).append(c)

        for semana, comps in sorted(por_semana.items(), key=lambda x: (x[0] is None, x[0])):
            for comp in comps:
                alertas_txt = " | ".join(f"{a.icono} {a.mensaje}" for a in comp.alertas)
                ws.append([
                    semana if semana is not None else "?",
                    f"Pago {comp.indice_pago}",
                    comp.tipo_documento,
                    comp.monto,
                    comp.rut_beneficiario,
                    comp.tipo_cuenta,
                    comp.numero_cuenta,
                    comp.banco,
                    comp.fecha_pago,
                    comp.estado_pago,
                    alertas_txt,
                ])
                fila = ws.max_row
                ws.cell(fila, 4).number_format = "#,##0"
                ws.cell(fila, 4).alignment = _ALIGN_R
                sev = _severidad_max(comp)
                if sev:
                    for col in range(1, ncols + 1):
                        ws.cell(fila, col).fill = _FILL[sev]

            # Subtotal si hay más de un pago en la semana
            if len(comps) > 1:
                subtotal = sum(c.monto for c in comps if c.monto)
                ws.append(["", "SUBTOTAL", "", subtotal] + [""] * (ncols - 4))
                fila = ws.max_row
                for col in range(1, ncols + 1):
                    ws.cell(fila, col).fill = _FILL["subtotal"]
                    ws.cell(fila, col).font = _SUB_FONT
                ws.cell(fila, 4).number_format = "#,##0"
                ws.cell(fila, 4).alignment = _ALIGN_R

        # Total categoría
        total_cat = sum(c.monto for c in comps_cat if c.monto)
        fill_total = "total_cui" if cat == "cuidados" else "total_sub"
        ws.append(["TOTAL", "", "", total_cat] + [""] * (ncols - 4))
        fila = ws.max_row
        for col in range(1, ncols + 1):
            ws.cell(fila, col).fill = _FILL[fill_total]
            ws.cell(fila, col).font = _TOT_FONT
        ws.cell(fila, 4).number_format = "#,##0"
        ws.cell(fila, 4).alignment = _ALIGN_R

    for i, w in enumerate(_WIDTHS_ALUMNO, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ── Helpers ───────────────────────────────────────────────────────────────────

def _nombre_archivo(resultado: ResultadoAuditoria) -> str:
    carpeta = re.sub(r"[^\w\-]", "_", resultado.carpeta)
    fecha = resultado.fecha.replace("-", "")
    return f"SUBSIDIOS_{carpeta}_{fecha}.xlsx"


def _nombre_hoja(nombre: str, usados: set[str]) -> str:
    """Nombre válido para hoja Excel: ≤31 chars, sin caracteres reservados, único."""
    sanitized = re.sub(r"[\[\]:*?/\\']", "", nombre).strip()[:28]
    if not sanitized:
        sanitized = "Alumno"
    candidate = sanitized
    n = 2
    while candidate.upper() in {u.upper() for u in usados}:
        candidate = f"{sanitized[:25]}_{n}"
        n += 1
    usados.add(candidate)
    return candidate


def _quote_sheet(name: str) -> str:
    """Envuelve el nombre de hoja en comillas simples si tiene espacios u otros chars."""
    if re.search(r"[ \t\[\]:*?/\\]", name):
        return f"'{name.replace(chr(39), chr(39) + chr(39))}'"
    return name


def _formato_alertas(errores: int, avisos: int) -> str:
    partes = []
    if errores:
        partes.append(f"🔴 {errores}")
    if avisos:
        partes.append(f"🟡 {avisos}")
    return "  ".join(partes)


def _contar_alertas(alumno: ResultadoAlumno) -> tuple[int, int]:
    """Devuelve (n_errores, n_avisos) excluyendo severidad 'info'."""
    todas = (
        list(alumno.alertas) +
        [a for c in alumno.comprobantes for a in c.alertas]
    )
    errores = sum(1 for a in todas if a.severidad == "error")
    avisos  = sum(1 for a in todas if a.severidad == "warn")
    return errores, avisos


def _severidad_max(c: Comprobante) -> str | None:
    if any(a.severidad == "error" for a in c.alertas):
        return "error"
    if any(a.severidad == "warn" for a in c.alertas):
        return "warn"
    return None
